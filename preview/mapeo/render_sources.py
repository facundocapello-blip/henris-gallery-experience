"""Render the supplied workbook's authentic header row; never rewrite its cells.

Usage: python3 render_sources.py /absolute/path/pedidos.xlsx
Dependencies: openpyxl, PyMuPDF, LibreOffice (available in the cloud environment).
Intermediate workbooks contain the original data and remain in a temporary folder.
Only the printed header row is retained as PNGs and metadata.
"""
import hashlib
import json
import pathlib
import subprocess
import sys
import tempfile

import fitz
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.page import PageMargins

ROOT = pathlib.Path(__file__).resolve().parent
GROUPS = [
    ('pedido', 'Pedido', 'A', 'C'),
    ('obras', 'Obras', 'D', 'H'),
    ('produccion', 'Producción', 'I', 'L'),
    ('cobro', 'Importe y moneda', 'M', 'P'),
    ('sena', 'Seña', 'Q', 'T'),
    ('final', 'Pago final', 'U', 'Y'),
    ('entrega', 'Entrega', 'Z', 'AA'),
    ('contabilidad', 'Contabilidad', 'AB', 'AF'),
    ('cliente', 'Cliente', 'AG', 'AJ'),
    ('facturacion', 'Facturación', 'AK', 'AL'),
    ('canal', 'Canal y web', 'AM', 'AN'),
]


def render(source):
    source = pathlib.Path(source).resolve()
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    original = openpyxl.load_workbook(source, data_only=False)
    sheet = original['Hoja1']
    headers = {get_column_letter(c): str(sheet.cell(1, c).value) for c in range(1, 41)}
    out = ROOT / 'sources'
    out.mkdir(exist_ok=True)
    manifest = dict(file=source.name, sha256=digest, sheet='Hoja1', row=1,
                    renderer='LibreOffice Calc → PDF → PNG (PyMuPDF)',
                    scope='Cabeceras originales únicamente; sin filas de clientes ni operaciones.',
                    groups=[])
    ranges = [('completo', 'Todas las cabeceras', 'A', 'AN')] + GROUPS
    with tempfile.TemporaryDirectory(prefix='henris-header-') as temp:
        temp = pathlib.Path(temp)
        inputs, pdfs = temp / 'input', temp / 'pdf'
        inputs.mkdir(); pdfs.mkdir()
        for key, _, first, last in ranges:
            book = openpyxl.load_workbook(source, data_only=False)
            s = book['Hoja1']
            s.print_area = f'{first}1:{last}1'
            s.print_options.headings = True
            s.sheet_properties.pageSetUpPr.fitToPage = True
            s.page_setup.paperSize = None
            s.page_setup.paperWidth = '5000pt'
            s.page_setup.paperHeight = '120pt'
            s.page_setup.orientation = 'portrait'
            s.page_setup.fitToWidth = 1
            s.page_setup.fitToHeight = 1
            s.page_margins = PageMargins(left=.15, right=.15, top=.12, bottom=.12, header=0, footer=0)
            for hf in [s.oddHeader, s.oddFooter, s.evenHeader, s.evenFooter, s.firstHeader, s.firstFooter]:
                for side in ['left', 'center', 'right']:
                    getattr(hf, side).text = ''
            book.save(inputs / f'{key}.xlsx')
        command = ['soffice', f'-env:UserInstallation={(temp / "profile").as_uri()}', '--headless', '--convert-to', 'pdf', '--outdir', str(pdfs)]
        subprocess.run(command + [str(x) for x in sorted(inputs.glob('*.xlsx'))], check=True, capture_output=True, text=True)
        for key, title, first, last in ranges:
            doc = fitz.open(pdfs / f'{key}.pdf')
            if len(doc) != 1:
                raise ValueError(f'{key}: expected one header page, got {len(doc)}')
            page = doc[0]
            drawings = page.get_drawings()
            bounds = fitz.Rect()
            for d in drawings:
                bounds |= d['rect']
            bounds = fitz.Rect(bounds.x0 - 4, bounds.y0 - 4, bounds.x1 + 4, bounds.y1 + 4)
            header_rects = sorted([d['rect'] for d in drawings if d['fill'] is None and 11 < d['rect'].height < 15 and d['rect'].width > 20], key=lambda r: r.x0)
            count = column_index_from_string(last) - column_index_from_string(first) + 1
            if len(header_rects) != count:
                raise ValueError(f'{key}: expected {count} column header rectangles, got {len(header_rects)}')
            image = page.get_pixmap(matrix=fitz.Matrix(3, 3), clip=bounds)
            image.save(out / f'{key}.png')
            cells = []
            for offset, r in enumerate(header_rects):
                letter = get_column_letter(column_index_from_string(first) + offset)
                cells.append(dict(column=letter, cell=f'{letter}1', label=headers[letter],
                                  x=(r.x0 - bounds.x0) / bounds.width,
                                  width=r.width / bounds.width))
            manifest['groups'].append(dict(id=key, title=title, range=f'{first}1:{last}1',
                                           image=f'sources/{key}.png', width=image.width, height=image.height,
                                           cells=cells))
    if hashlib.sha256(source.read_bytes()).hexdigest() != digest:
        raise RuntimeError('Original workbook unexpectedly changed')
    (out / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')
    print(f'{len(ranges)} captures rendered. Original unchanged: {digest}')


if __name__ == '__main__':
    render(sys.argv[1])
